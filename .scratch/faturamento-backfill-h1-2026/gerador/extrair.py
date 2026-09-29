# Extrai as linhas das abas mensais das planilhas Q1/Q2 para JSON cru.
import openpyxl, json, datetime, warnings, sys
warnings.filterwarnings('ignore')
COLS = ['convenio','data_fat','responsavel','protocolo','lote','data_prov','valor_enviado','valor_recebido',
        'data_receb','status','saldo_pendente','glosas','glosa_devida','refaturamento','lote_refat','valor_refat',
        'recurso1','recurso2','valor_acatado','saldo_receber']
def cv(v):
    if isinstance(v,(datetime.datetime,datetime.date)): return {'d':v.strftime('%Y-%m-%d')}
    if isinstance(v,datetime.time): return {'t':str(v)}
    return v
out=[]
for f,abas in [('q1.xlsx',['JANEIRO','FEVEREIRO','MARÇO']),('q2.xlsx',['ABRIL','MAIO','JUNHO'])]:
    wb=openpyxl.load_workbook(f,data_only=True)
    for aba in abas:
        ws=wb[aba]
        hdr=[c for c in next(ws.iter_rows(min_row=24,max_row=24,values_only=True))[2:22]]
        assert hdr[0]=='Convênio' and hdr[19]=='Saldo a Receber', (aba,hdr)
        for i,r in enumerate(ws.iter_rows(min_row=25,values_only=True),25):
            vals=r[2:22]
            if all(v is None or v=='' for v in vals): continue
            out.append({'aba':aba,'linha':i,**{k:cv(v) for k,v in zip(COLS,vals)}})
json.dump(out,open('h1/linhas.json','w'),ensure_ascii=False,indent=0)
print(len(out))
